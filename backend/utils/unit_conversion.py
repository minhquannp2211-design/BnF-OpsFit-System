from functools import lru_cache
from math import isfinite
from numbers import Real
from pathlib import Path
import re

import openpyxl


WORKBOOK_PATH = (
    Path(__file__).resolve().parent.parent.parent/ "data"/ "templates"/ "BnF_Opsfit_System.xlsx")


def _unit_key(unit: object) -> str:
    if not isinstance(unit, str) or not unit.strip():
        raise ValueError("Unit must be a non-empty string")
    return "".join(unit.casefold().split())
@lru_cache(maxsize=1)
def _unit_index() -> dict[str, tuple[str, str, float]]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        sheet = workbook["units"]
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        fields = {
            match.group(1).strip().casefold()
            for value in headers
            if value is not None
            for match in [re.search(r"\(([^()]*)\)\s*$", str(value))]
            if match
        }
        required_fields = {
            "unit_code", "unit_name", "unit_group", "target_unit",
            "conversion_factor", "aliases",
        }
        if missing := required_fields - fields:
            raise ValueError(f"Missing columns {sorted(missing)} in sheet 'units'")
        definitions: dict[str, tuple[str, str, float]] = {}
        alias_rows: list[tuple[str, object]] = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None for value in row):
                continue
            code = _unit_key(row[0])
            group = _unit_key(row[2])
            standard_unit = _unit_key(row[3])
            try:
                factor = float(row[4])
            except (TypeError, ValueError) as error:
                raise ValueError(
                    f"Invalid conversion factor in units row {row_number}"
                ) from error
            if code in definitions:
                raise ValueError(f"Duplicated unit '{code}'")
            if not isfinite(factor) or factor <= 0:
                raise ValueError(f"Invalid factor for unit '{code}'")
            definitions[code] = (group, standard_unit, factor)
            alias_rows.append((code, row[5] if len(row) > 5 else None))

        for code, (group, standard_unit, _) in definitions.items():
            target = definitions.get(standard_unit)
            if target is None or target[0] != group or target[2] != 1:
                raise ValueError(
                    f"Invalid target unit '{standard_unit}' for '{code}'"
                )

        unit_index = definitions.copy()
        for code, raw_aliases in alias_rows:
            if not raw_aliases:
                continue
            for alias in str(raw_aliases).split("|"):
                alias_key = _unit_key(alias)
                definition = definitions[code]
                if (
                    alias_key in unit_index
                    and unit_index[alias_key] != definition
                ):
                    raise ValueError(f"Duplicated unit alias '{alias}'")
                unit_index[alias_key] = definition

        return unit_index
    finally:
        workbook.close()


def _unit_definition(unit: str) -> tuple[str, str, float]:
    definition = _unit_index().get(_unit_key(unit))
    if definition is None:
        raise ValueError(f"Unsupported unit: '{unit}'")
    return definition


def target_unit(unit: str) -> str:
    return _unit_definition(unit)[1]     #Return the target unit configured for unit in Excel.


def unit_group(unit: str) -> str:
    return _unit_definition(unit)[0]


def are_compatible_units(first_unit: str, second_unit: str) -> bool:
    return unit_group(first_unit) == unit_group(second_unit)


def to_standard_quantity(quantity: Real, unit: str) -> float:
    if isinstance(quantity, bool) or not isinstance(quantity, Real):
        raise ValueError("Quantity must be a number")

    numeric_quantity = float(quantity)
    if not isfinite(numeric_quantity) or numeric_quantity < 0:
        raise ValueError("Quantity must be a finite, non-negative number")

    conversion_factor = _unit_definition(unit)[2]
    return numeric_quantity * conversion_factor


def standardize_quantity(
    quantity: Real,
    unit: str,
    expected_unit: str | None = None,
) -> tuple[float, str]:
    """Return a quantity in its standard unit, checking an optional reference unit."""
    standard_unit = target_unit(unit)
    if expected_unit is not None and not are_compatible_units(
        unit, expected_unit
    ):
        raise ValueError(
            f"Incompatible units '{unit}' and '{expected_unit}'"
        )
    return to_standard_quantity(quantity, unit), standard_unit
