from collections import defaultdict
from pathlib import Path
import openpyxl

from typing import TypeAlias
import math
import re
import unicodedata #for creating food_id


CURRENT_DIR = Path(__file__).resolve().parent
FILE_PATH = CURRENT_DIR.parent.parent / "data" / "templates" / "BnF_Opsfit_System.xlsx"
workbook = openpyxl.load_workbook(FILE_PATH, read_only = True, data_only = True)
menu_sheet = workbook["menu"]
recipe_sheet = workbook["recipe"]
ingredients_sheet = workbook["ingredients"]

Food: TypeAlias = tuple[
    str,  #food_id
    str,  #food_name
    str,  #food_size
    float #food_price
]
MenuKey: TypeAlias = tuple[str,str]
menu: list[Food] = []
menu_key_index: dict[MenuKey, str] = {}
menu_id_index: dict[str,Food] = {}


def normalize_text(value: str) -> str:
    value = value.replace("Đ", "D").replace("đ", "d")
    value = unicodedata.normalize("NFD", value)
    value = "".join(
        character
        for character in value
        if unicodedata.category(character) != "Mn"
    )
    value = value.upper().strip()
    value = re.sub(r"[^A-Z0-9]+", "_", value)
    return value.strip("_")

def require_text(
    value: object,
    field_name: str,
    sheet_name: str,
    row_number: int,
) -> str:
    if value is None:
        raise ValueError(
            f"Missing '{field_name}' in sheet "
            f"'{sheet_name}', row {row_number}"
        )

    text = str(value).strip()

    if not text or not normalize_text(text):
        raise ValueError(
            f"Invalid '{field_name}' value '{text}' "
            f"in sheet '{sheet_name}', row {row_number}"
        )

    return text

def require_number(
    value: object,
    field_name: str,
    sheet_name: str,
    row_number: int,
) -> float:
    if value is None:
        raise ValueError(
            f"Missing '{field_name}' in sheet "
            f"'{sheet_name}', row {row_number}"
        )

    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(
            f"Invalid numeric value '{value}' for "
            f"'{field_name}' in sheet '{sheet_name}', "
            f"row {row_number}"
        ) from error

    if not math.isfinite(number):
        raise ValueError(
            f"'{field_name}' must be a finite number "
            f"in sheet '{sheet_name}', row {row_number}"
        )

    return number

def validate_optional_id(
    value: object,
    expected_id: str,
    field_name: str,
    sheet_name: str,
    row_number: int,
) -> None:
    if value is None or not str(value).strip():
        return

    provided_id = normalize_text(str(value))

    if provided_id != expected_id:
        raise ValueError(
            f"Expected '{field_name}' '{expected_id}', "
            f"received '{provided_id}' in sheet "
            f"'{sheet_name}', row {row_number}"
        )

def create_food_id(food_name: str, food_size: str) -> str:
    normalized_name = normalize_text(food_name)
    normalized_size = normalize_text(food_size)

    return f"{normalized_name}_{normalized_size}"
def create_ingredient_id(ingredient_name: str) -> str:
    normalized_name = normalize_text(ingredient_name)

    return f"ING_{normalized_name}"
for row_number, row in enumerate(
    menu_sheet.iter_rows(min_row=2, values_only=True),
    start=2,
):
    if not any(value is not None for value in row):
        continue
    food_name = require_text(
        row[1], "food_name", "menu", row_number
    )
    food_size = require_text(
        row[2], "food_size", "menu", row_number
    )
    food_id = create_food_id(food_name, food_size)
    validate_optional_id(
        row[0], food_id, "food_id", "menu", row_number
    )
    food_price = require_number(
        row[3], "food_price", "menu", row_number
    )
    if food_price < 0:
        raise ValueError(
            f"Food '{food_name}' has invalid price "
            f"'{food_price}'"
        )
    food = (food_id, food_name, food_size, food_price)
    menu_key: MenuKey = (
        normalize_text(food_name),
        normalize_text(food_size),
    )
    if menu_key in menu_key_index:
        raise ValueError(
            f"Duplicated food '{food_name}' "
            f"size '{food_size}' in menu sheet"
        )

    if food_id in menu_id_index:
        raise ValueError(
            f"Duplicated food_id '{food_id}' "
            f"in menu sheet"
        )
    menu.append(food)
    menu_key_index[menu_key] = food_id
    menu_id_index[food_id] = food

RecipeKey: TypeAlias = tuple[str, str, str]  # recipe_id, food_name, food_size
RecipeDetail: TypeAlias = tuple[
    str,    # food_id
    str,    # ingredient_id
    str,    # ingredient_name
    float,  # ingredient_quantity
    str,    # ingredient_unit
]
RecipeDict: TypeAlias = dict[RecipeKey,list[RecipeDetail]]

recipe_id_index: dict[str, RecipeKey] = {}
recipe_food_id_index: dict[str, RecipeKey] = {}
recipe_name_index: dict[str,list[RecipeKey]] = defaultdict(list)
recipe_name_size_index: dict[MenuKey, RecipeKey] = {}
recipe_ingredient_index: set[tuple[str, str]] = set()
recipe_dict: RecipeDict = defaultdict(list)

for row_number, row in enumerate(
    recipe_sheet.iter_rows(min_row=2, values_only=True),
    start=2,
):
    if not any(value is not None for value in row):
        continue
    raw_food_name = require_text(
        row[2], "food_name", "recipe", row_number
    )
    raw_food_size = require_text(
        row[4], "food_size", "recipe", row_number
    )
    menu_key: MenuKey = (
        normalize_text(raw_food_name),
        normalize_text(raw_food_size),
    )
    food_id = menu_key_index.get(menu_key)
    if food_id is None:
        raise ValueError(
            f"Cannot find '{raw_food_name}' size "
            f"'{raw_food_size}' in menu sheet, "
            f"recipe row {row_number}"
        )

    food = menu_id_index[food_id]
    food_name = food[1]
    food_size = food[2]
    recipe_id = f"RCP_{food_id}"
    validate_optional_id(
        row[0], recipe_id, "recipe_id", "recipe", row_number
    )
    validate_optional_id(
        row[1], food_id, "food_id", "recipe", row_number
    )
    ingredient_name = require_text(
        row[5], "ingredient_name", "recipe", row_number
    )
    ingredient_id = create_ingredient_id(ingredient_name)
    validate_optional_id(
        row[3], ingredient_id, "ingredient_id", "recipe", row_number
    )
    ingredient_quantity = require_number(
        row[6], "ingredient_quantity", "recipe", row_number
    )
    if ingredient_quantity <= 0:
        raise ValueError(
            f"Ingredient '{ingredient_name}' in "
            f"recipe '{recipe_id}' must have "
            f"quantity greater than zero"
        )
    ingredient_unit = require_text(
        row[7], "ingredient_unit", "recipe", row_number
    )

    recipe_key: RecipeKey = (
        recipe_id,
        food_name,
        food_size,
    )
    recipe_detail: RecipeDetail = (
        food_id,
        ingredient_id,
        ingredient_name,
        ingredient_quantity,
        ingredient_unit,
    )

    recipe_ingredient_key = (
        recipe_id,
        ingredient_id,
    )

    if recipe_ingredient_key in recipe_ingredient_index:
        raise ValueError(
            f"Duplicated ingredient '{ingredient_name}' "
            f"in recipe '{recipe_id}'"
        )

    existing_key = recipe_id_index.get(recipe_id)

    if existing_key is not None and existing_key != recipe_key:
        raise ValueError(
            f"Recipe ID '{recipe_id}' is associated "
            f"with multiple recipe keys"
        )

    existing_key = recipe_food_id_index.get(food_id)

    if existing_key is not None and existing_key != recipe_key:
        raise ValueError(
            f"Food ID '{food_id}' has multiple recipes"
        )

    recipe_ingredient_index.add(recipe_ingredient_key)
    recipe_dict[recipe_key].append(recipe_detail)
    recipe_id_index[recipe_id] = recipe_key
    recipe_food_id_index[food_id] = recipe_key
    recipe_name_size_index[menu_key] = recipe_key

    normalized_name = normalize_text(food_name)

    if recipe_key not in recipe_name_index[normalized_name]:
        recipe_name_index[normalized_name].append(recipe_key)

def find_recipe(
    search_value: str,
    food_size: str | None = None,
) -> tuple[RecipeKey, list[RecipeDetail]] | None:
    search_value = search_value.strip()
    normalized_value = normalize_text(search_value)
    recipe_key = recipe_id_index.get(normalized_value)

    if recipe_key is None:
        recipe_key = recipe_food_id_index.get(normalized_value)

    if recipe_key is None and food_size is not None:
        menu_key: MenuKey = (
            normalized_value,
            normalize_text(food_size),
        )

        recipe_key = recipe_name_size_index.get(menu_key)

    if recipe_key is not None:
        return recipe_key, recipe_dict[recipe_key]

    matched_keys = recipe_name_index.get(normalized_value, [])

    if not matched_keys:
        return None

    if len(matched_keys) > 1:
        available_sizes = ", ".join(
            matched_key[2]
            for matched_key in matched_keys
        )

        raise ValueError(
            f"Food '{search_value}' has multiple sizes: "
            f"{available_sizes}. Please specify food_size."
        )

    recipe_key = matched_keys[0]

    return recipe_key, recipe_dict[recipe_key]

Ingredient: TypeAlias = tuple [
    str,      #ingredient_id
    str,      #ingredient_name
    float,    #package_price
    float,    #package_quantity
    str,      #package_unit ( don vi tinh theo g,ml, ... )
    float,    #waste_rate
]
ingredients: list[Ingredient] = []
ingredient_id_index: dict[str, Ingredient] = {}
ingredient_name_index: dict[str, str] = {}
for row_number, row in enumerate(
    ingredients_sheet.iter_rows(min_row=2, values_only=True),
    start=2,
):
    if not any(value is not None for value in row):
        continue
    ingredient_name = require_text(
        row[1], "ingredient_name", "ingredients", row_number
    )
    ingredient_id = create_ingredient_id(ingredient_name)
    validate_optional_id(
        row[0], ingredient_id,
        "ingredient_id", "ingredients", row_number,
    )
    package_price = require_number(
        row[2], "package_price", "ingredients", row_number
    )
    package_quantity = require_number(
        row[3], "package_quantity", "ingredients", row_number
    )
    package_unit = require_text(
        row[4], "package_unit", "ingredients", row_number
    )
    raw_waste_rate = require_number(
        row[5], "waste_rate", "ingredients", row_number
    )
    waste_rate = (
        raw_waste_rate / 100
        if raw_waste_rate > 1
        else raw_waste_rate
    )

    if ingredient_id in ingredient_id_index:
        raise ValueError(
            f"Duplicated ingredient "
            f"'{ingredient_name}' in ingredients sheet"
        )

    if package_price < 0:
        raise ValueError(
            f"Ingredient '{ingredient_name}' "
            f"has invalid package price"
        )

    if package_quantity <= 0:
        raise ValueError(
            f"Ingredient '{ingredient_name}' "
            f"must have package quantity greater than zero"
        )

    if not 0 <= waste_rate <= 1:
        raise ValueError(
            f"Ingredient '{ingredient_name}' "
            f"has invalid waste rate '{waste_rate}'"
        )

    ingredient: Ingredient = (
        ingredient_id,
        ingredient_name,
        package_price,
        package_quantity,
        package_unit,
        waste_rate,
    )

    ingredients.append(ingredient)
    ingredient_id_index[ingredient_id] = ingredient
    ingredient_name_index[
        normalize_text(ingredient_name)
    ] = ingredient_id
for recipe_key, recipe_details in recipe_dict.items():
    for recipe_detail in recipe_details:
        food_id = recipe_detail[0]
        ingredient_id = recipe_detail[1]
        ingredient_name = recipe_detail[2]
        ingredient_unit = recipe_detail[4]

        ingredient = ingredient_id_index.get(ingredient_id)

        if ingredient is None:
            raise ValueError(
                f"Recipe '{recipe_key[0]}' uses "
                f"unknown ingredient "
                f"'{ingredient_name}'"
            )

        package_unit = ingredient[4]

        if (
            normalize_text(ingredient_unit)
            != normalize_text(package_unit)
        ):
            raise ValueError(
                f"Unit mismatch for ingredient "
                f"'{ingredient_name}' in recipe "
                f"'{recipe_key[0]}': recipe uses "
                f"'{ingredient_unit}', ingredients "
                f"sheet uses '{package_unit}'"
            )

workbook.close()
